#!/usr/bin/env python3
"""Background worker threads for the MIoT GUI."""

import json
import os
import time
import traceback

import requests
from PyQt6.QtCore import QThread, pyqtSignal

from miot_export_template import (
    parse_prop_row,
    parse_action_row,
    parse_event_row,
    write_prop_sheet,
    write_action_sheet,
    write_event_sheet,
    write_config_sheet,
    write_source_sheet,
)
from miot_create_properties import (
    match_service,
    build_request_body,
    build_action_request_body,
    build_event_request_body,
    create_property,
    create_action,
    create_event,
    HEADERS,
    QUERY_SERVICES_API,
)
from miot_service_core import (
    get_services,
    sync_services,
    parse_service_str,
    modify_iid,
)
from miot_automation_core import (
    get_automation_list,
    sync_automations,
    write_automation_export_excel,
)
from miot_common import (
    parse_json_response,
    is_success_response,
    response_message,
)
from miot_reports import write_execution_report, default_desktop_path


class ExportPropWorker(QThread):
    """导出属性模板"""
    progress    = pyqtSignal(str)
    finished_ok  = pyqtSignal(str)
    finished_err = pyqtSignal(str)

    def __init__(self, pid, model, token, ph, userid, connect_type,
                 output_path, save_json, delay):
        super().__init__()
        self.pid = pid; self.model = model; self.token = token
        self.ph = ph; self.userid = userid; self.connect_type = connect_type
        self.output_path = output_path; self.save_json = save_json
        self.delay = delay
        self._cancel_flag = False

    def cancel(self):
        self._cancel_flag = True

    def run(self):
        try:
            cfg = {
                "userId": str(self.userid), "pdId": str(self.pid),
                "model": self.model, "serviceToken": self.token,
                "xiaomiiot_ph": self.ph,
                "connectType": str(self.connect_type),
            }
            cookies = {"serviceToken": self.token, "userId": str(self.userid),
                       "xiaomiiot_ph": self.ph}
            params_base = {"userId": str(self.userid),
                           "xiaomiiot_ph": self.ph, "pdId": str(self.pid)}

            all_props = []
            all_actions = []
            all_events = []
            max_retries = 3

            def _safe_get(url, params, label, retries=max_retries):
                """带重试的 GET 请求，返回 result 列表"""
                for attempt in range(1, retries + 1):
                    try:
                        r = requests.get(url, params=params, headers=HEADERS, cookies=cookies, timeout=15)
                        data = parse_json_response(r, f"{label}查询 API")
                        return data.get("result", []) if is_success_response(data) else []
                    except Exception as e:
                        if attempt < retries:
                            self.progress.emit(f"  ⚠️ {label}查询失败(第{attempt}次)，1s后重试: {e}")
                            time.sleep(1)
                        else:
                            self.progress.emit(f"  ❌ {label}查询失败(已重试{retries}次): {e}")
                            raise

            self.progress.emit("📋 正在查询产品服务列表...")
            params = {**params_base, "model": self.model,
                      "connectType": str(self.connect_type),
                      "language": "zh_cn", "version": "1", "status": "0"}
            services = None
            for attempt in range(1, max_retries + 1):
                try:
                    resp = requests.get(QUERY_SERVICES_API, params=params,
                                        headers=HEADERS, cookies=cookies, timeout=15)
                    if resp.status_code != 200:
                        self.finished_err.emit(f"HTTP 请求失败 (status={resp.status_code})\n请检查网络连接")
                        return
                    try:
                        data = parse_json_response(resp, "查询服务列表 API")
                    except RuntimeError as e:
                        self.finished_err.emit(str(e))
                        return
                    if not is_success_response(data):
                        self.finished_err.emit(f"查询服务失败: {response_message(data, str(data))}")
                        return
                    services = data.get("result", [])
                    break
                except Exception as e:
                    if attempt < max_retries:
                        self.progress.emit(f"  ⚠️ 服务列表查询失败(第{attempt}次)，1s后重试: {e}")
                        time.sleep(1)
                    else:
                        self.finished_err.emit(f"❌ 服务列表查询失败(已重试{max_retries}次): {e}")
                        return
            if not services:
                self.finished_err.emit("未查到服务，请检查 Cookie 和产品信息")
                return
            self.progress.emit(f"✅ 找到 {len(services)} 个服务")

            for i, svc in enumerate(services):
                if self._cancel_flag:
                    self.finished_err.emit("⚠️ 用户取消")
                    return
                siid = svc.get("siid", "?")
                sname = svc.get("description", svc.get("name", ""))
                stype = svc.get("type", "")
                self.progress.emit(f"🔍 [{i+1}/{len(services)}] 查询服务 siid={siid} ({sname})...")
                params2 = {**params_base, "version": "1", "status": "0",
                           "siid": str(siid), "serviceType": stype,
                           "model": self.model,
                           "connectType": str(self.connect_type), "language": "zh_cn"}

                # 属性
                props = _safe_get(
                    "https://iot.mi.com/cgi-std/api/v1/functionDefine/getInstanceProperties",
                    params2, "属性")
                for p in props:
                    p["_service"] = svc
                all_props.extend(props)

                # 方法
                actions = _safe_get(
                    "https://iot.mi.com/cgi-std/api/v1/functionDefine/getInstanceActions",
                    params2, "方法")
                for a in actions:
                    a["_service"] = svc
                all_actions.extend(actions)

                # 事件
                events = _safe_get(
                    "https://iot.mi.com/cgi-std/api/v1/functionDefine/getInstanceEvents",
                    params2, "事件")
                for e in events:
                    e["_service"] = svc
                all_events.extend(events)

                if self.delay > 0:
                    time.sleep(self.delay)

            self.progress.emit(f"✅ 共获取 {len(all_props)} 条属性, {len(all_actions)} 个方法, {len(all_events)} 个事件")

            if not self.output_path:
                safe_model = self.model.replace(".", "_").replace("-", "_")
                self.output_path = os.path.join(os.path.expanduser("~"), "Desktop", f"MIoT_模板_{safe_model}.xlsx")

            self.progress.emit("📝 正在生成 Excel 模板...")
            from openpyxl import Workbook
            wb = Workbook()
            ws1 = wb.active; ws1.title = "属性定义"
            rows_data = [parse_prop_row(p, p.get("_service", {})) for p in all_props]
            write_prop_sheet(ws1, rows_data)

            # 方法定义 Sheet
            action_rows = [parse_action_row(a, a.get("_service", {})) for a in all_actions]
            ws2 = wb.create_sheet("方法定义")
            write_action_sheet(ws2, action_rows)

            # 事件定义 Sheet
            event_rows = [parse_event_row(e, e.get("_service", {})) for e in all_events]
            ws3 = wb.create_sheet("事件定义")
            write_event_sheet(ws3, event_rows)

            # 公共配置 Sheet
            ws4 = wb.create_sheet("公共配置")

            class _Args:
                pass
            args = _Args()
            args.pid = self.pid; args.model = self.model
            args.token = self.token; args.ph = self.ph
            args.userid = self.userid; args.connect_type = self.connect_type
            write_config_sheet(ws4, args)

            ws5 = wb.create_sheet("原始数据参考")
            write_source_sheet(ws5, services, rows_data, action_rows, event_rows)
            wb.save(self.output_path)

            if self.save_json:
                json_path = self.output_path.replace(".xlsx", ".json")
                with open(json_path, "w", encoding="utf-8") as f:
                    json.dump({"services": services, "properties": all_props,
                               "actions": all_actions, "events": all_events},
                              f, ensure_ascii=False, indent=2, default=str)
                self.progress.emit(f"💾 JSON 已保存: {json_path}")

            self.finished_ok.emit(self.output_path)
        except Exception:
            self.finished_err.emit(f"导出失败:\n{traceback.format_exc()}")


class CreatePropWorker(QThread):
    """批量创建属性"""
    progress         = pyqtSignal(str)
    update_progress  = pyqtSignal(int, int)
    finished_ok      = pyqtSignal(int, int)
    finished_err     = pyqtSignal(str)

    def __init__(self, config, props, services, delay):
        super().__init__()
        self.config = config; self.props = props
        self.services = services; self.delay = delay
        self._cancel_flag = False

    def cancel(self):
        self._cancel_flag = True

    def run(self):
        try:
            success = failed = 0
            results = []
            for i, prop in enumerate(self.props):
                if self._cancel_flag:
                    self.progress.emit("⚠️ 用户取消，已停止创建")
                    break
                name = prop.get("name", f"行{i+1}")
                svc = match_service(prop, self.services)
                siid = svc["siid"] if svc else prop.get("siid", "?")
                sname = svc.get("description", svc.get("name", "")) if svc else "未匹配"
                body = build_request_body(prop, self.config, svc)
                self.update_progress.emit(i + 1, len(self.props))
                self.progress.emit(f"  [{i+1}] {name} → siid={siid} ({sname}) ...")
                try:
                    resp = create_property(body, self.config)
                    if is_success_response(resp):
                        piid = resp.get("result")
                        # 校验并修正 piid
                        expected_piid = prop.get("piid")
                        if expected_piid and str(expected_piid).strip():
                            try:
                                expected_piid_int = int(expected_piid)
                                if int(piid) != expected_piid_int:
                                    self.progress.emit(f"    🔧 PIID {piid}→{expected_piid_int} 修正中...")
                                    r = modify_iid(self.config, siid, piid, expected_piid_int, "PIID")
                                    if is_success_response(r):
                                        self.progress.emit(f"  ✅ {name} 成功 (piid={expected_piid_int}, 已修正)")
                                        success += 1
                                        results.append({"name": name, "status": "success", "piid": expected_piid_int, "original_piid": piid, "siid": siid})
                                    else:
                                        msg_m = response_message(r, str(r))
                                        self.progress.emit(f"  ⚠️ {name} 修正失败: {msg_m}")
                                        self.progress.emit(f"  ✅ {name} 成功 (piid={piid})")
                                        success += 1
                                        results.append({"name": name, "status": "success", "piid": piid, "modify_error": msg_m, "siid": siid})
                                else:
                                    self.progress.emit(f"  ✅ {name} 成功 (piid={piid})")
                                    success += 1
                                    results.append({"name": name, "status": "success", "piid": piid, "siid": siid})
                            except (ValueError, TypeError):
                                self.progress.emit(f"  ✅ {name} 成功 (piid={piid})")
                                success += 1
                                results.append({"name": name, "status": "success", "piid": piid, "siid": siid})
                        else:
                            self.progress.emit(f"  ✅ {name} 成功 (piid={piid})")
                            success += 1
                            results.append({"name": name, "status": "success", "piid": piid, "siid": siid})
                    else:
                        msg = response_message(resp, json.dumps(resp, ensure_ascii=False))
                        self.progress.emit(f"  ❌ {name} 失败 ({msg})")
                        failed += 1
                        results.append({"name": name, "status": "failed", "error": msg, "siid": siid})
                except Exception as e:
                    self.progress.emit(f"  ❌ {name} 异常 ({e})")
                    failed += 1
                    results.append({"name": name, "status": "error", "error": str(e), "siid": siid})
                if self.delay > 0:
                    time.sleep(self.delay)

            result_path = os.path.join(os.path.expanduser("~"), "Desktop", "miot_create_result.json")
            with open(result_path, "w", encoding="utf-8") as f:
                json.dump(results, f, ensure_ascii=False, indent=2)
            report_path = default_desktop_path("miot_create_report.xlsx")
            write_execution_report(report_path, [{"type": "属性", **row} for row in results])
            self.progress.emit(f"📊 执行报告已保存: {report_path}")
            self.finished_ok.emit(success, failed)
        except Exception:
            self.finished_err.emit(f"创建失败:\n{traceback.format_exc()}")


class CreateAllWorker(QThread):
    """批量创建属性+方法+事件"""
    progress         = pyqtSignal(str)
    update_progress  = pyqtSignal(int, int)
    finished_ok      = pyqtSignal(int, int)
    finished_err     = pyqtSignal(str)

    def __init__(self, config, tasks, services, delay):
        """
        tasks: [(type_label, item, build_fn, create_fn, id_field, svc), ...]
        """
        super().__init__()
        self.config = config; self.tasks = tasks
        self.services = services; self.delay = delay
        self._cancel_flag = False

    def cancel(self):
        self._cancel_flag = True

    def run(self):
        try:
            success = failed = 0
            results = []
            for i, (type_label, item, build_fn, create_fn, id_field, svc) in enumerate(self.tasks):
                if self._cancel_flag:
                    self.progress.emit("⚠️ 用户取消，已停止创建")
                    break
                name = item.get("name", f"行{i+1}")
                siid = svc["siid"] if svc else item.get("siid", "?")
                sname = svc.get("description", svc.get("name", "")) if svc else "未匹配"
                body = build_fn(item, self.config, svc)
                self.update_progress.emit(i + 1, len(self.tasks))
                self.progress.emit(f"  [{i+1}][{type_label}] {name} → siid={siid} ({sname}) ...")
                try:
                    resp = create_fn(body, self.config)
                    if is_success_response(resp):
                        new_id = resp.get("result")
                        # 校验并修正 ID
                        expected_id = item.get(id_field)
                        which_iid_map = {"piid": "PIID", "aiid": "AIID", "eiid": "EIID"}
                        which_iid = which_iid_map.get(id_field, "")
                        if expected_id and which_iid and str(expected_id).strip():
                            try:
                                expected_id_int = int(expected_id)
                                if int(new_id) != expected_id_int:
                                    self.progress.emit(f"    🔧 {which_iid} {new_id}→{expected_id_int} 修正中...")
                                    r = modify_iid(self.config, siid, new_id, expected_id_int, which_iid)
                                    if is_success_response(r):
                                        self.progress.emit(f"  ✅ [{type_label}] {name} 成功 ({id_field}={expected_id_int}, 已修正)")
                                        success += 1
                                        results.append({"type": type_label, "name": name, "status": "success", id_field: expected_id_int, "original_id": new_id, "siid": siid})
                                    else:
                                        msg_m = response_message(r, str(r))
                                        self.progress.emit(f"  ⚠️ [{type_label}] {name} 修正失败: {msg_m}")
                                        self.progress.emit(f"  ✅ [{type_label}] {name} 成功 ({id_field}={new_id})")
                                        success += 1
                                        results.append({"type": type_label, "name": name, "status": "success", id_field: new_id, "modify_error": msg_m, "siid": siid})
                                else:
                                    self.progress.emit(f"  ✅ [{type_label}] {name} 成功 ({id_field}={new_id})")
                                    success += 1
                                    results.append({"type": type_label, "name": name, "status": "success", id_field: new_id, "siid": siid})
                            except (ValueError, TypeError):
                                self.progress.emit(f"  ✅ [{type_label}] {name} 成功 ({id_field}={new_id})")
                                success += 1
                                results.append({"type": type_label, "name": name, "status": "success", id_field: new_id, "siid": siid})
                        else:
                            self.progress.emit(f"  ✅ [{type_label}] {name} 成功 ({id_field}={new_id})")
                            success += 1
                            results.append({"type": type_label, "name": name, "status": "success", id_field: new_id, "siid": siid})
                    else:
                        msg = response_message(resp, json.dumps(resp, ensure_ascii=False))
                        self.progress.emit(f"  ❌ [{type_label}] {name} 失败 ({msg})")
                        failed += 1
                        results.append({"type": type_label, "name": name, "status": "failed", "error": msg, "siid": siid})
                except Exception as e:
                    self.progress.emit(f"  ❌ [{type_label}] {name} 异常 ({e})")
                    failed += 1
                    results.append({"type": type_label, "name": name, "status": "error", "error": str(e), "siid": siid})
                if self.delay > 0:
                    time.sleep(self.delay)

            result_path = os.path.join(os.path.expanduser("~"), "Desktop", "miot_create_result.json")
            with open(result_path, "w", encoding="utf-8") as f:
                json.dump(results, f, ensure_ascii=False, indent=2)
            report_path = default_desktop_path("miot_create_report.xlsx")
            write_execution_report(report_path, results)
            self.progress.emit(f"📊 执行报告已保存: {report_path}")
            self.finished_ok.emit(success, failed)
        except Exception:
            self.finished_err.emit(f"创建失败:\n{traceback.format_exc()}")


class SyncServiceWorker(QThread):
    """批量同步服务（创建 / 修正 siid）"""
    progress    = pyqtSignal(str)
    finished_ok  = pyqtSignal(dict)
    finished_err = pyqtSignal(str)

    def __init__(self, config, service_rows, dry_run, delay=0.5):
        super().__init__()
        self.config = config; self.service_rows = service_rows
        self.dry_run = dry_run; self._cancel = False
        self.delay = delay

    def cancel(self):
        self._cancel = True

    def run(self):
        try:
            result = sync_services(
                self.config, self.service_rows,
                dry_run=self.dry_run,
                log_fn=self.progress.emit,
                cancelled_fn=lambda: self._cancel,
            )
            results_path = os.path.join(os.path.expanduser("~"), "Desktop", "sync_results.json")
            with open(results_path, "w", encoding="utf-8") as f:
                json.dump(result["results"], f, ensure_ascii=False, indent=2)
            self.finished_ok.emit(result)
        except Exception:
            self.finished_err.emit(f"同步失败:\n{traceback.format_exc()}")


# ─── 自动化 Worker ────────────────────────────────────────────

class ExportAutomationWorker(QThread):
    """导出自动化列表"""
    progress    = pyqtSignal(str)
    finished_ok  = pyqtSignal(str)
    finished_err = pyqtSignal(str)

    def __init__(self, config, output_path):
        super().__init__()
        self.config = config; self.output_path = output_path

    def run(self):
        try:
            self.progress.emit("📋 正在查询自动化列表...")
            auto_list = get_automation_list(self.config)
            then_count = sum(1 for a in auto_list if a.get("_trType") == "then")
            if_count = sum(1 for a in auto_list if a.get("_trType") == "if")
            then_action = sum(1 for a in auto_list if a.get("_trType") == "then" and a.get("actionList"))
            then_simple = then_count - then_action
            self.progress.emit(f"✅ 找到 {len(auto_list)} 个自动化（执行动作: {then_count}[组合{then_action}+普通{then_simple}], 触发条件: {if_count}）")

            if not auto_list:
                self.finished_err.emit("未查到自动化，请检查 Cookie 和产品信息")
                return

            self.progress.emit("📝 正在生成 Excel...")
            write_automation_export_excel(self.output_path, self.config, auto_list)
            self.finished_ok.emit(self.output_path)
        except Exception:
            self.finished_err.emit(f"导出失败:\n{traceback.format_exc()}")


class CreateAutomationWorker(QThread):
    """批量创建自定义自动化"""
    progress    = pyqtSignal(str)
    update_progress = pyqtSignal(int, int)  # current, total
    finished_ok  = pyqtSignal(int, int)     # success, failed
    finished_err = pyqtSignal(str)

    def __init__(self, config, auto_items, dry_run=False, delay=0.5):
        super().__init__()
        self.config = config; self.auto_items = auto_items
        self.dry_run = dry_run; self.delay = delay
        self._cancel = False

    def cancel(self):
        self._cancel = True

    def run(self):
        try:
            result = sync_automations(
                self.config, self.auto_items,
                dry_run=self.dry_run,
                delay=self.delay,
                log_fn=self.progress.emit,
                cancelled_fn=lambda: self._cancel,
            )
            s = len(result["success"])
            f = len(result["failed"])
            # 保存结果
            results_path = os.path.join(os.path.expanduser("~"), "Desktop", "automation_results.json")
            with open(results_path, "w", encoding="utf-8") as fout:
                json.dump(result, fout, ensure_ascii=False, indent=2, default=str)
            self.finished_ok.emit(s, f)
        except Exception:
            self.finished_err.emit(f"创建失败:\n{traceback.format_exc()}")


class ExportServiceWorker(QThread):
    """导出服务 / 属性详情"""
    progress    = pyqtSignal(str)
    finished_ok  = pyqtSignal(str)
    finished_err = pyqtSignal(str)

    def __init__(self, config, output_path, export_props=False):
        super().__init__()
        self.config = config; self.output_path = output_path
        self.export_props = export_props

    def run(self):
        try:
            import pandas as pd
            self.progress.emit(f"正在获取 {self.config.get('model')} 的服务列表...")
            services = get_services(self.config)
            self.progress.emit(f"✅ 获取到 {len(services)} 个服务")

            config_rows = [
                {"参数名": "userId",       "值": self.config.get("userId", "")},
                {"参数名": "xiaomiiot_ph", "值": self.config.get("xiaomiiot_ph", "")},
                {"参数名": "serviceToken", "值": self.config.get("serviceToken", "")},
                {"参数名": "pdId",         "值": self.config.get("pdId", "")},
                {"参数名": "model",        "值": self.config.get("model", "")},
            ]
            df_config = pd.DataFrame(config_rows)

            svc_rows = []
            prop_rows = []
            for svc in services:
                siid  = svc.get("siid", "")
                sname = svc.get("name", "")
                sdesc = svc.get("description", "") or svc.get("normalizationDesc", "")
                ndesc = svc.get("normalizationDesc", "")
                std   = "true" if svc.get("standard") else "false"
                svc_rows.append({
                    "服务ID": siid, "服务名称": sname, "服务描述": sdesc,
                    "标准化描述": ndesc, "是否标准服务": std,
                })
                if self.export_props:
                    parsed = parse_service_str(svc)
                    for prop in parsed["properties"]:
                        ptype = prop.get("type", "")
                        pname = ptype.split(":")[-2] if ":" in ptype else ""
                        prop_rows.append({
                            "siid": siid, "服务名称": sname,
                            "piid": prop.get("iid", ""), "类型": "属性",
                            "属性名称": pname, "描述": prop.get("description", ""),
                            "格式": prop.get("format", ""),
                            "访问权限": ",".join(prop.get("access", [])),
                            "值列表": json.dumps(prop.get("value-list", []), ensure_ascii=False) if prop.get("value-list") else "",
                            "值范围": json.dumps(prop.get("value-range", []), ensure_ascii=False) if prop.get("value-range") else "",
                        })
                    for evt in parsed["events"]:
                        etype = evt.get("type", "")
                        prop_rows.append({
                            "siid": siid, "服务名称": sname,
                            "piid": evt.get("iid", ""), "类型": "事件",
                            "属性名称": etype.split(":")[-2] if ":" in etype else "",
                            "描述": evt.get("description", ""),
                            "格式": "", "访问权限": "", "值列表": "", "值范围": "",
                        })
                    for act in parsed["actions"]:
                        atype = act.get("type", "")
                        prop_rows.append({
                            "siid": siid, "服务名称": sname,
                            "piid": act.get("iid", ""), "类型": "动作",
                            "属性名称": atype.split(":")[-2] if ":" in atype else "",
                            "描述": act.get("description", ""),
                            "格式": "", "访问权限": "", "值列表": "", "值范围": "",
                        })
                    self.progress.emit(
                        f"  siid={siid} {sname}  属性:{len(parsed['properties'])} "
                        f"事件:{len(parsed['events'])} 动作:{len(parsed['actions'])}"
                    )

            df_svc = pd.DataFrame(svc_rows)
            with pd.ExcelWriter(self.output_path, engine="openpyxl") as writer:
                df_config.to_excel(writer, index=False, sheet_name="产品配置")
                df_svc.to_excel(writer, index=False, sheet_name="服务列表")
                if self.export_props and prop_rows:
                    pd.DataFrame(prop_rows).to_excel(writer, index=False, sheet_name="属性详情")

            self.finished_ok.emit(self.output_path)
        except Exception:
            self.finished_err.emit(f"导出失败:\n{traceback.format_exc()}")

