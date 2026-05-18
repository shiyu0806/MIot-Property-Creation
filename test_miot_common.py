#!/usr/bin/env python3
"""
miot_common.py + miot_create_properties.py 核心函数单元测试

运行: python3 -m pytest test_miot_common.py -v
或:   python3 test_miot_common.py
"""

import sys
import os
import unittest
import tempfile
from unittest.mock import Mock, patch

from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from miot_common import (
    build_cookies, build_params, build_headers, safe_int, safe_request,
    parse_json_response, is_success_response, response_message,
    BASE, DEFAULT_HEADERS, TEMPLATE_VERSION,
    PROPERTY_COLUMNS, ACTION_COLUMNS, EVENT_COLUMNS,
)
from miot_create_properties import (
    match_service, detect_value_type, parse_value_list,
    build_request_body, build_action_request_body, build_event_request_body,
    _build_action_event_base, parse_bool, parse_access,
    load_property_excel, read_properties,
    validate_config, validate_items, validate_tasks,
)
from miot_reports import write_dry_run_plan, write_execution_report


# ═══════════════════════════════════════════════════════════════════
# miot_common 测试
# ═══════════════════════════════════════════════════════════════════

class TestSafeInt(unittest.TestCase):
    def test_normal_int(self):
        assert safe_int("123") == 123

    def test_none(self):
        assert safe_int(None) == 0

    def test_empty_string(self):
        assert safe_int("") == 0

    def test_whitespace(self):
        assert safe_int("   ") == 0

    def test_invalid(self):
        assert safe_int("abc") == 0

    def test_custom_default(self):
        assert safe_int(None, default=-1) == -1

    def test_float_string(self):
        assert safe_int("3.14") == 0

    def test_negative(self):
        assert safe_int("-5") == -5


class TestBuildCookies(unittest.TestCase):
    def test_normal(self):
        cfg = {"serviceToken": "abc", "userId": "123", "xiaomiiot_ph": "xyz"}
        cookies = build_cookies(cfg)
        assert cookies["serviceToken"] == "abc"
        assert cookies["userId"] == "123"
        assert cookies["xiaomiiot_ph"] == "xyz"

    def test_missing_keys(self):
        cookies = build_cookies({})
        assert cookies["serviceToken"] == ""
        assert cookies["userId"] == ""

    def test_numeric_values(self):
        cookies = build_cookies({"userId": 12345})
        assert cookies["userId"] == "12345"


class TestBuildParams(unittest.TestCase):
    def test_normal(self):
        cfg = {"userId": "u1", "xiaomiiot_ph": "p1", "extra": "ignored"}
        params = build_params(cfg)
        assert params == {"userId": "u1", "xiaomiiot_ph": "p1"}


class TestBuildHeaders(unittest.TestCase):
    def test_default(self):
        h = build_headers()
        assert h["referer"] == f"{BASE}/"
        assert "user-agent" in h

    def test_custom_referer(self):
        h = build_headers("https://example.com")
        assert h["referer"] == "https://example.com"


class TestApiResponseHelpers(unittest.TestCase):
    def test_parse_json_response_ok(self):
        resp = Mock(status_code=200)
        resp.json.return_value = {"status": 200, "result": []}
        assert parse_json_response(resp, "测试 API") == {"status": 200, "result": []}

    def test_parse_json_response_non_json_login_hint(self):
        resp = Mock(status_code=200, text="<html>passport login</html>")
        resp.json.side_effect = ValueError("no json")
        try:
            parse_json_response(resp, "测试 API")
        except RuntimeError as exc:
            msg = str(exc)
            assert "测试 API 返回非 JSON" in msg
            assert "登录态失效" in msg
        else:
            raise AssertionError("parse_json_response should reject non-json responses")

    def test_parse_json_response_rejects_list(self):
        resp = Mock(status_code=200)
        resp.json.return_value = []
        try:
            parse_json_response(resp, "测试 API")
        except RuntimeError as exc:
            assert "返回结构异常" in str(exc)
        else:
            raise AssertionError("parse_json_response should reject non-object JSON")

    def test_is_success_response(self):
        assert is_success_response({"status": 200}) is True
        assert is_success_response({"code": 0}) is True
        assert is_success_response({"status": 500}) is False

    def test_response_message(self):
        assert response_message({"message": "失败"}) == "失败"
        assert response_message({"msg": "错误"}) == "错误"
        assert response_message({}) == "未知错误"


class TestConstants(unittest.TestCase):
    def test_base_url(self):
        assert BASE == "https://iot.mi.com"

    def test_column_definitions(self):
        for col_def in PROPERTY_COLUMNS + ACTION_COLUMNS + EVENT_COLUMNS:
            assert len(col_def) == 5, f"列定义应为 5 元组: {col_def}"
            assert isinstance(col_def[2], int), f"宽度应为 int: {col_def}"
            assert isinstance(col_def[4], bool), f"required 应为 bool: {col_def}"

    def test_template_version(self):
        assert TEMPLATE_VERSION


class TestExcelLoading(unittest.TestCase):
    def test_read_properties_uses_name_column(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["name", "description", "format"])
        ws.append(["说明", "说明", "说明"])
        ws.append([None, "空名称", "bool"])
        ws.append(["on", "开关", "bool"])
        rows = read_properties(ws)
        assert len(rows) == 1
        assert rows[0]["name"] == "on"

    def test_load_property_excel_missing_config_sheet(self):
        wb = Workbook()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            path = tmp.name
        try:
            wb.save(path)
            try:
                load_property_excel(path)
            except ValueError as exc:
                assert "公共配置" in str(exc)
            else:
                raise AssertionError("load_property_excel should require 公共配置 Sheet")
        finally:
            os.remove(path)


class TestReportWriters(unittest.TestCase):
    def test_write_dry_run_plan(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            path = tmp.name
        try:
            write_dry_run_plan(path, [{
                "type": "属性", "index": 1, "name": "on",
                "plan_status": "待创建",
            }])
            assert os.path.exists(path)
            assert os.path.getsize(path) > 0
        finally:
            os.remove(path)

    def test_write_execution_report(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            path = tmp.name
        try:
            write_execution_report(path, [{
                "type": "属性", "name": "on", "status": "success", "piid": 1,
            }])
            assert os.path.exists(path)
            assert os.path.getsize(path) > 0
        finally:
            os.remove(path)

    def test_load_property_excel_ok(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "属性定义"
        ws.append(["name", "description", "format"])
        ws.append(["说明", "说明", "说明"])
        ws.append(["on", "开关", "bool"])
        ws2 = wb.create_sheet("公共配置")
        ws2.append(["配置项", "值", "说明"])
        ws2.append(["template_version", TEMPLATE_VERSION, ""])
        ws2.append(["userId", "123", ""])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            path = tmp.name
        try:
            wb.save(path)
            config, props, actions, events = load_property_excel(path)
            assert config["template_version"] == TEMPLATE_VERSION
            assert config["userId"] == "123"
            assert props[0]["name"] == "on"
            assert actions == []
            assert events == []
        finally:
            os.remove(path)


# ═══════════════════════════════════════════════════════════════════
# 输入校验测试
# ═══════════════════════════════════════════════════════════════════

class TestValidation(unittest.TestCase):
    def test_validate_config_ok(self):
        cfg = {
            "serviceToken": "token", "xiaomiiot_ph": "ph",
            "userId": "123", "pdId": "456", "model": "test.device",
        }
        assert validate_config(cfg) == []

    def test_validate_config_missing(self):
        errors = validate_config({"userId": "123"})
        assert any("serviceToken" in e for e in errors)
        assert any("pdId" in e for e in errors)

    def test_validate_property_items_ok(self):
        items = [{
            "name": "mode", "description": "模式", "format": "uint8",
            "value_list": "0:关闭,1:开启",
        }]
        assert validate_items(items, "property", "属性定义") == []

    def test_validate_property_items_rejects_bad_enum(self):
        items = [{
            "name": "mode", "description": "模式", "format": "uint8",
            "value_list": "broken",
        }]
        errors = validate_items(items, "property", "属性定义")
        assert any("value_list 格式错误" in e for e in errors)

    def test_validate_property_items_rejects_bad_format(self):
        items = [{"name": "x", "description": "X", "format": "object"}]
        errors = validate_items(items, "property", "属性定义")
        assert any("format 不支持" in e for e in errors)

    def test_validate_property_items_accepts_float_range(self):
        items = [{
            "name": "temperature", "description": "温度", "format": "float",
            "value_range_min": "-20.5", "value_range_max": "80.5",
            "value_range_step": "0.1",
        }]
        assert validate_items(items, "property", "属性定义") == []

    def test_validate_property_items_rejects_bad_float_range(self):
        items = [{
            "name": "temperature", "description": "温度", "format": "float",
            "value_range_min": "low",
        }]
        errors = validate_items(items, "property", "属性定义")
        assert any("不是有效数字" in e for e in errors)

    def test_validate_property_items_allows_duplicate_name(self):
        items = [
            {"name": "mode", "description": "模式1", "format": "uint8"},
            {"name": "mode", "description": "模式2", "format": "uint8"},
        ]
        assert validate_items(items, "property", "属性定义") == []

    def test_validate_items_duplicate_name(self):
        items = [
            {"name": "toggle", "description": "切换", "service_desc": "开关一键"},
            {"name": "toggle", "description": "切换2", "service_desc": "开关一键"},
        ]
        errors = validate_items(items, "action", "方法定义")
        assert any("name 重复" in e for e in errors)

    def test_validate_items_allows_duplicate_name_across_services(self):
        items = [
            {"name": "toggle", "description": "切换1", "service_desc": "开关一键"},
            {"name": "toggle", "description": "切换2", "service_desc": "开关二键"},
            {"name": "toggle", "description": "切换3", "service_desc": "开关三键"},
        ]
        assert validate_items(items, "action", "方法定义") == []

    def test_validate_tasks_rejects_missing_siid(self):
        tasks = [{"index": 1, "name": "on", "siid": "?"}]
        errors = validate_tasks(tasks, "属性定义")
        assert any("未匹配到有效服务" in e for e in errors)


# ═══════════════════════════════════════════════════════════════════
# match_service 测试
# ═══════════════════════════════════════════════════════════════════

class TestMatchService(unittest.TestCase):
    def setUp(self):
        self.services = [
            {"siid": 2, "name": "light", "type": "urn:miot-spec-v2:device:light:1", "description": "Light"},
            {"siid": 3, "name": "switch", "type": "urn:miot-spec-v2:service:switch:1", "description": "Switch"},
            {"siid": 4, "name": "mode", "type": "urn:miot-spec-v2:service:mode:1", "description": "Mode"},
            {"siid": 10, "name": "dimmer", "type": "urn:miot-spec-v2:service:dimmer:1", "description": "Dimmer"},
            {"siid": 20, "name": "countdown", "type": "urn:miot-spec-v2:service:countdown:1", "description": "Countdown"},
        ]

    def test_exact_desc_and_name(self):
        prop = {"service_desc": "Light", "name": "on", "service_name": "light"}
        result = match_service(prop, self.services)
        assert result is not None
        assert result["siid"] == 2

    def test_exact_desc_only(self):
        prop = {"service_desc": "Switch", "service_name": "xxx"}
        result = match_service(prop, self.services)
        assert result is not None
        assert result["siid"] == 3

    def test_fuzzy_desc_includes(self):
        """优先级 5: service_desc 在 description 中"""
        prop = {"service_desc": "ight", "service_name": "xxx"}
        result = match_service(prop, self.services)
        assert result is not None  # "ight" 是 "Light" 的子串
        assert result["siid"] == 2

    def test_name_matching(self):
        prop = {"service_desc": "nonexistent", "service_name": "countdown"}
        result = match_service(prop, self.services)
        assert result is not None
        assert result["siid"] == 20

    def test_siid_fallback(self):
        prop = {"siid": 4, "service_desc": "nonexistent", "service_name": "xxx"}
        result = match_service(prop, self.services)
        assert result is not None
        assert result["siid"] == 4

    def test_no_match(self):
        prop = {"service_desc": "nonexistent", "service_name": "xxx"}
        result = match_service(prop, self.services)
        assert result is None

    def test_empty_services(self):
        prop = {"service_desc": "Light"}
        result = match_service(prop, [])
        assert result is None

    def test_exact_desc_priority_over_siid(self):
        prop = {"siid": 99, "service_desc": "Switch", "service_name": "switch"}
        result = match_service(prop, self.services)
        assert result is not None
        assert result["siid"] == 3


# ═══════════════════════════════════════════════════════════════════
# detect_value_type / parse_value_list 测试
# ═══════════════════════════════════════════════════════════════════

class TestDetectValueType(unittest.TestCase):
    def test_enum(self):
        prop = {"format": "uint8", "value_list": "0:关闭,1:开启,2:待机", "value_range_min": "", "value_range_max": ""}
        assert detect_value_type("uint8", prop) == "enum"

    def test_number(self):
        prop = {"format": "uint16", "value_list": "", "value_range_min": "0", "value_range_max": "100"}
        assert detect_value_type("uint16", prop) == "number"

    def test_bool_type(self):
        prop = {"format": "bool", "value_list": "", "value_range_min": "", "value_range_max": ""}
        assert detect_value_type("bool", prop) == "bool_range"

    def test_string_type(self):
        prop = {"format": "string", "value_list": "", "value_range_min": "", "value_range_max": ""}
        assert detect_value_type("string", prop) == "string"


class TestParseValueList(unittest.TestCase):
    def test_standard(self):
        result = parse_value_list("0:关闭,1:开启,2:待机")
        assert result == [
            {"description": "关闭", "value": 0},
            {"description": "开启", "value": 1},
            {"description": "待机", "value": 2},
        ]

    def test_whitespace(self):
        result = parse_value_list(" 0 : 关闭 , 1 : 开启 ")
        assert len(result) == 2
        assert result[0]["value"] == 0

    def test_empty(self):
        assert parse_value_list("") == []
        assert parse_value_list(None) == []

    def test_single(self):
        result = parse_value_list("1:on")
        assert result == [{"description": "on", "value": 1}]


# ═══════════════════════════════════════════════════════════════════
# parse_bool / parse_access 测试
# ═══════════════════════════════════════════════════════════════════

class TestParseBool(unittest.TestCase):
    def test_true_values(self):
        for v in ["1", "true", "True", "TRUE", "yes", 1, True]:
            assert parse_bool(v) is True, f"应为 True: {v!r}"

    def test_false_values(self):
        for v in ["0", "false", "no", "off", "", None, 0]:
            assert parse_bool(v) is False, f"应为 False: {v!r}"


class TestParseAccess(unittest.TestCase):
    def test_standard(self):
        assert parse_access("read,write,notify") == ["read", "write", "notify"]

    def test_single(self):
        assert parse_access("read") == ["read"]

    def test_empty(self):
        assert parse_access("") == ["read", "write", "notify"]

    def test_whitespace(self):
        assert parse_access(" read , write ") == ["read", "write"]


# ═══════════════════════════════════════════════════════════════════
# build_request_body 测试
# ═══════════════════════════════════════════════════════════════════

class TestBuildRequestBody(unittest.TestCase):
    def setUp(self):
        self.config = {"model": "test.device", "pdId": "123", "version": "1", "language": "zh_cn"}
        self.service_info = {"siid": 3, "type": "urn:test:switch:1"}

    def test_basic_property(self):
        prop = {
            "name": "on", "description": "开关", "format": "bool",
            "access": "read,write,notify",
            "value_list": "", "value_range_min": "", "value_range_max": "",
            "value_range_step": "", "standard": "", "valueType": "", "unit": "", "piid": "",
        }
        body = build_request_body(prop, self.config, self.service_info)
        assert body["siid"] == 3
        assert body["description"] == "开关"
        assert body["format"] == "bool"
        assert body["access"] == ["read", "write", "notify"]

    def test_enum_property(self):
        prop = {
            "name": "mode", "description": "模式", "format": "uint8",
            "access": "read,write", "value_list": "0:自动,1:手动",
            "value_range_min": "", "value_range_max": "", "value_range_step": "",
            "standard": "", "valueType": "", "unit": "", "piid": "",
        }
        body = build_request_body(prop, self.config, self.service_info)
        assert body["format"] == "uint8"
        assert body["access"] == ["read", "write"]

    def test_number_property(self):
        prop = {
            "name": "brightness", "description": "亮度", "format": "uint8",
            "access": "read,write", "value_list": "",
            "value_range_min": "0", "value_range_max": "100", "value_range_step": "1",
            "standard": "", "valueType": "", "unit": "", "piid": "",
        }
        body = build_request_body(prop, self.config, self.service_info)
        assert body["format"] == "uint8"
        assert body["valueRange"] == [0, 100, 1]

    def test_float_property_range(self):
        prop = {
            "name": "temperature", "description": "温度", "format": "float",
            "access": "read,notify", "value_list": "",
            "value_range_min": "-20.5", "value_range_max": "80.5", "value_range_step": "0.1",
            "standard": "", "valueType": "", "unit": "", "piid": "",
        }
        body = build_request_body(prop, self.config, self.service_info)
        assert body["format"] == "float"
        assert body["valueRange"] == [-20.5, 80.5, 0.1]


# ═══════════════════════════════════════════════════════════════════
# build_action / build_event 测试
# ═══════════════════════════════════════════════════════════════════

class TestBuildActionEventBody(unittest.TestCase):
    def setUp(self):
        self.config = {"model": "test.device", "pdId": "123", "version": "1", "language": "zh_cn"}
        self.service_info = {"siid": 3, "type": "urn:test:switch:1"}

    def test_action_body(self):
        item = {"name": "toggle", "description": "切换", "normalizationDesc": "toggle"}
        body = build_action_request_body(item, self.config, self.service_info)
        assert body["siid"] == 3
        assert body["name"] == "toggle"
        assert body["description"] == "切换"

    def test_event_body(self):
        item = {"name": "alarm", "description": "告警", "normalizationDesc": "alarm"}
        body = build_event_request_body(item, self.config, self.service_info)
        assert body["siid"] == 3
        assert body["name"] == "alarm"

    def test_action_defaults(self):
        item = {"name": "reset", "description": "重置", "normalizationDesc": ""}
        body = build_action_request_body(item, self.config, self.service_info)
        assert body["normalizationDesc"] == "reset"


# ═══════════════════════════════════════════════════════════════════
# safe_request 测试
# ═══════════════════════════════════════════════════════════════════

class TestSafeRequest(unittest.TestCase):
    def test_success(self):
        mocked_resp = Mock(status_code=200)
        with patch("requests.request", return_value=mocked_resp) as request:
            resp = safe_request("GET", "https://example.com/get", timeout=5, max_retries=1)
        assert resp.status_code == 200
        request.assert_called_once()

    def test_retry_on_connection_error(self):
        with patch("requests.request", side_effect=ConnectionError("boom")) as request:
            try:
                safe_request(
                    "GET", "http://localhost:1", timeout=1,
                    max_retries=2, retry_delay=0, log_fn=lambda _: None,
                )
            except Exception as exc:
                assert "boom" in str(exc)
            else:
                raise AssertionError("safe_request should raise after retries")
        assert request.call_count == 2


if __name__ == "__main__":
    print("运行 miot_common + miot_create_properties 核心函数单元测试...\n")
    unittest.main(verbosity=2)
