#!/usr/bin/env python3
"""Excel report helpers for dry-run plans and execution results."""

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def default_desktop_path(filename: str) -> str:
    """Return a Desktop path, falling back to the home directory if needed."""
    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    return os.path.join(desktop if os.path.isdir(desktop) else os.path.expanduser("~"), filename)


def _write_table(path: str, title: str, rows: list[dict], columns: list[tuple[str, str]]):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    header_row = 4
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, (_, header) in enumerate(columns, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, header_row + 1):
        for col_idx, (key, _) in enumerate(columns, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{max(header_row, header_row + len(rows))}"
    for col_idx, (key, header) in enumerate(columns, 1):
        max_len = len(header)
        for row in rows:
            max_len = max(max_len, len(str(row.get(key, ""))))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 45)

    wb.save(path)
    return path


def write_dry_run_plan(path: str, rows: list[dict]):
    columns = [
        ("type", "类型"),
        ("index", "行号"),
        ("name", "name"),
        ("description", "description"),
        ("format", "format"),
        ("value_type", "值类型"),
        ("siid", "siid"),
        ("service", "服务"),
        ("plan_status", "计划状态"),
        ("note", "备注"),
    ]
    return _write_table(path, "MIoT dry-run 计划表", rows, columns)


def write_execution_report(path: str, rows: list[dict]):
    columns = [
        ("type", "类型"),
        ("name", "name"),
        ("status", "状态"),
        ("siid", "siid"),
        ("piid", "piid"),
        ("aiid", "aiid"),
        ("eiid", "eiid"),
        ("original_id", "原 ID"),
        ("original_piid", "原 piid"),
        ("modified", "已修正"),
        ("modify_error", "修正错误"),
        ("error", "错误"),
    ]
    return _write_table(path, "MIoT 执行报告", rows, columns)
