#!/usr/bin/env python3
"""Property management tabs for the MIoT GUI."""

import os
import traceback

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QFormLayout, QLineEdit, QPushButton,
    QCheckBox, QFileDialog, QSpinBox, QGroupBox, QMessageBox,
)

from miot_create_properties import (
    query_services as create_query_services,
    match_service,
    build_request_body,
    build_action_request_body,
    build_event_request_body,
    create_property,
    create_action,
    create_event,
    detect_value_type,
    load_property_excel,
    validate_config,
    validate_items,
    validate_tasks,
)
from miot_service_core import check_product_status
from miot_common import TEMPLATE_VERSION
from miot_gui_common import (
    _make_log_panel,
    _make_progress,
    _inject_group_id,
    _cookie_group,
)
from miot_gui_workers import ExportPropWorker, CreateAllWorker
from miot_reports import write_dry_run_plan, default_desktop_path


class ExportPropTab(QWidget):
    def __init__(self):
        super().__init__()
        self._worker = None
        self._build()

    def _build(self):
        layout = QHBoxLayout(self)
        left = QWidget(); left.setFixedWidth(460)
        lv = QVBoxLayout(left)

        # 产品信息
        grp_prod = QGroupBox("产品信息")
        form = QFormLayout()
        self.pid = QLineEdit(); self.pid.setPlaceholderText("如 33257")
        self.model = QLineEdit(); self.model.setPlaceholderText("如 uwize.switch.yzw07")
        self.userid = QLineEdit(); self.userid.setPlaceholderText("如 1097752639")
        self.connect_type = QSpinBox()
        self.connect_type.setRange(0, 99); self.connect_type.setValue(16)
        form.addRow("产品ID (pdId):", self.pid)
        form.addRow("产品型号 (model):", self.model)
        form.addRow("用户ID (userId):", self.userid)
        form.addRow("连接类型:", self.connect_type)
        grp_prod.setLayout(form)
        lv.addWidget(grp_prod)

        _, self.token, self.ph, _ = _cookie_group(lv, "exp_prop", show_userid=False)

        # 输出
        grp_out = QGroupBox("输出选项")
        form2 = QFormLayout()
        self.out_edit = QLineEdit(); self.out_edit.setPlaceholderText("点击浏览选择导出文件夹")
        self.out_edit.setReadOnly(True)
        btn_br = QPushButton("浏览...")
        btn_br.clicked.connect(self._browse_out)
        row = QHBoxLayout(); row.addWidget(self.out_edit); row.addWidget(btn_br)
        form2.addRow("导出文件夹:", row)
        self.chk_json = QCheckBox("同时保存原始 JSON")
        form2.addRow("", self.chk_json)
        grp_out.setLayout(form2)
        lv.addWidget(grp_out)

        btn_row = QHBoxLayout()
        self.btn_start = QPushButton("🚀 开始导出")
        self.btn_start.setObjectName("successBtn")
        self.btn_start.clicked.connect(self._start)
        self.btn_cancel = QPushButton("取消")
        self.btn_cancel.clicked.connect(self._cancel)
        self.btn_cancel.setEnabled(False)
        btn_row.addWidget(self.btn_start); btn_row.addWidget(self.btn_cancel)
        lv.addLayout(btn_row)
        lv.addStretch()

        right = QWidget(); rv = QVBoxLayout(right)
        self.log = _make_log_panel(rv)
        self.progress = _make_progress(rv)

        layout.addWidget(left); layout.addWidget(right, stretch=1)

    def _browse_out(self):
        path = QFileDialog.getExistingDirectory(self, "选择导出文件夹")
        if path:
            self.out_edit.setText(path)

    def _start(self):
        pid = self.pid.text().strip()
        model = self.model.text().strip()
        token = self.token.text().strip()
        ph = self.ph.text().strip()
        userid = self.userid.text().strip()
        if not all([pid, model, token, ph, userid]):
            QMessageBox.warning(self, "提示", "请填写产品信息和 Cookie")
            return

        # 自动生成输出路径
        out_dir = self.out_edit.text().strip()
        if not out_dir:
            out_dir = os.path.join(os.path.expanduser("~"), "Desktop")
        safe_model = model.replace(".", "_").replace("-", "_")
        output_path = os.path.join(out_dir, f"MIoT_模板_{safe_model}.xlsx")

        self.log.clear()
        self.btn_start.setEnabled(False); self.btn_cancel.setEnabled(True)
        self.progress.setVisible(True); self.progress.setRange(0, 0)

        self._worker = ExportPropWorker(
            pid, model, token, ph, userid,
            self.connect_type.value(),
            output_path,
            self.chk_json.isChecked(),
            0,  # 导出不需要间隔
        )
        self._worker.progress.connect(self.log.append)
        self._worker.finished_ok.connect(self._done_ok)
        self._worker.finished_err.connect(self._done_err)
        self._worker.start()

    def _cancel(self):
        if self._worker and self._worker.isRunning():
            self._worker.cancel()
            self.log.append("⚠️ 取消请求已发送")
        self._reset()

    def _done_ok(self, path):
        self._reset()
        self.log.append(f"\n🎉 导出成功: {path}")
        QMessageBox.information(self, "导出成功", f"模板已保存:\n{path}")

    def _done_err(self, msg):
        self._reset()
        self.log.append(f"\n❌ {msg}")
        QMessageBox.critical(self, "导出失败", msg)

    def _reset(self):
        self.btn_start.setEnabled(True); self.btn_cancel.setEnabled(False)
        self.progress.setVisible(False)


# ─── Tab: 创建属性 ────────────────────────────────────────────

class CreatePropTab(QWidget):
    def __init__(self):
        super().__init__()
        self._worker = None
        self._build()

    def _build(self):
        layout = QHBoxLayout(self)
        left = QWidget(); left.setFixedWidth(460)
        lv = QVBoxLayout(left)

        grp_file = QGroupBox("Excel 文件")
        fv = QHBoxLayout()
        self.file_edit = QLineEdit(); self.file_edit.setPlaceholderText("选择属性 Excel 文件")
        btn_br = QPushButton("选择文件")
        btn_br.clicked.connect(self._browse_file)
        fv.addWidget(self.file_edit); fv.addWidget(btn_br)
        grp_file.setLayout(fv)
        lv.addWidget(grp_file)

        grp_ov = QGroupBox("产品信息（可覆盖 Excel 配置）")
        form_ov = QFormLayout()
        self.pid_ov = QLineEdit(); self.pid_ov.setPlaceholderText("留空使用 Excel 配置")
        self.model_ov = QLineEdit(); self.model_ov.setPlaceholderText("留空使用 Excel 配置")
        form_ov.addRow("产品ID (pdId):", self.pid_ov)
        form_ov.addRow("产品型号 (model):", self.model_ov)
        grp_ov.setLayout(form_ov)
        lv.addWidget(grp_ov)

        _, self.token_ov, self.ph_ov, self.uid_ov = _cookie_group(lv, "crt_prop")
        self.token_ov.setPlaceholderText("留空使用 Excel 配置")
        self.ph_ov.setPlaceholderText("留空使用 Excel 配置")
        self.uid_ov.setPlaceholderText("留空使用 Excel 配置")

        grp_opts = QGroupBox("选项")
        form_opts = QFormLayout()
        self.delay_spin = QSpinBox()
        self.delay_spin.setRange(100, 2000); self.delay_spin.setValue(500)
        self.delay_spin.setSingleStep(100); self.delay_spin.setSuffix(" ms")
        self.siid_spin = QSpinBox()
        self.siid_spin.setRange(0, 999); self.siid_spin.setValue(0)
        self.siid_spin.setSpecialValueText("全部")
        form_opts.addRow("请求间隔:", self.delay_spin)
        form_opts.addRow("指定 siid:", self.siid_spin)
        grp_opts.setLayout(form_opts)
        lv.addWidget(grp_opts)

        btn_row1 = QHBoxLayout()
        self.btn_dry = QPushButton("🧪 干跑检查")
        self.btn_dry.clicked.connect(self._dryrun)
        self.btn_list = QPushButton("📋 查看服务")
        self.btn_list.clicked.connect(self._list_services)
        btn_row1.addWidget(self.btn_dry); btn_row1.addWidget(self.btn_list)
        lv.addLayout(btn_row1)

        btn_row2 = QHBoxLayout()
        self.btn_create = QPushButton("🚀 开始创建")
        self.btn_create.setObjectName("dangerBtn")
        self.btn_create.clicked.connect(self._start_create)
        self.btn_cancel = QPushButton("取消")
        self.btn_cancel.clicked.connect(self._cancel)
        self.btn_cancel.setEnabled(False)
        btn_row2.addWidget(self.btn_create); btn_row2.addWidget(self.btn_cancel)
        lv.addLayout(btn_row2)
        lv.addStretch()

        right = QWidget(); rv = QVBoxLayout(right)
        self.log = _make_log_panel(rv)
        self.progress = _make_progress(rv)

        layout.addWidget(left); layout.addWidget(right, stretch=1)

    def _browse_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "选择属性 Excel", "", "Excel (*.xlsx *.xls)")
        if path:
            self.file_edit.setText(path)

    def _load(self):
        path = self.file_edit.text().strip()
        if not path or not os.path.exists(path):
            QMessageBox.warning(self, "提示", "请选择有效的 Excel 文件")
            return None, None
        try:
            config, props, actions, events = load_property_excel(path)
        except (FileNotFoundError, ValueError) as e:
            QMessageBox.warning(self, "Excel 读取失败", str(e))
            return None, None

        for key, widget in [
            ("serviceToken", self.token_ov), ("xiaomiiot_ph", self.ph_ov),
            ("userId", self.uid_ov), ("pdId", self.pid_ov), ("model", self.model_ov),
        ]:
            if widget.text().strip():
                config[key] = widget.text().strip()

        if not props and not actions and not events:
            QMessageBox.warning(self, "提示", "属性/方法/事件定义均为空")
            return None, None

        validation_errors = []
        validation_errors.extend(validate_config(config))
        validation_errors.extend(validate_items(props, "property", "属性定义"))
        validation_errors.extend(validate_items(actions, "action", "方法定义"))
        validation_errors.extend(validate_items(events, "event", "事件定义"))
        if validation_errors:
            QMessageBox.warning(self, "Excel 校验未通过", "\n".join(validation_errors[:20]))
            return None, None

        # 自动注入 groupId
        _inject_group_id(config)

        return config, (props, actions, events)

    def _list_services(self):
        config, _ = self._load()
        if not config:
            return
        self.log.clear(); self.log.append("📋 查询服务列表...")
        try:
            services = create_query_services(config)
            if not services:
                self.log.append("❌ 未查到服务"); return
            self.log.append(f"\n📋 共 {len(services)} 个服务:")
            self.log.append("-" * 60)
            for svc in services:
                self.log.append(
                    f"{svc.get('siid','?'):>4} | {svc.get('name','?'):<24} | "
                    f"{svc.get('description',''):<20} | {svc.get('type','')}")
        except Exception as e:
            self.log.append(f"❌ {e}")

    def _dryrun(self):
        config, items = self._load()
        if not config:
            return
        props, actions, events = items
        self.log.clear(); self.log.append("🧪 干跑模式...\n")
        try:
            services = create_query_services(config)
            target_siid = self.siid_spin.value()
            plan_rows = []

            # 属性任务
            self.log.append("📝 属性定义:")
            prop_tasks = []
            for i, prop in enumerate(props):
                svc = match_service(prop, services)
                siid = svc["siid"] if svc else prop.get("siid", "?")
                if target_siid > 0 and str(siid) != str(target_siid):
                    continue
                sname = svc.get("description", svc.get("name", "")) if svc else "❌ 未匹配"
                vtype = detect_value_type(str(prop.get("format", "")), prop)
                prop_tasks.append((i+1, prop.get("name","?"), prop.get("format","?"), vtype, siid, sname))
                plan_rows.append({
                    "type": "属性",
                    "index": i + 1,
                    "name": prop.get("name", ""),
                    "description": prop.get("description", ""),
                    "format": prop.get("format", ""),
                    "value_type": vtype,
                    "siid": siid,
                    "service": sname,
                    "plan_status": "待创建" if svc else "服务未匹配",
                    "note": "dry-run 未执行",
                })

            if prop_tasks:
                self.log.append(f"{'#':>3} | {'name':<20} | {'format':<8} | {'vtype':<10} | siid | 服务")
                self.log.append("-" * 80)
                for t in prop_tasks:
                    self.log.append(f"{t[0]:>3} | {t[1]:<20} | {t[2]:<8} | {t[3]:<10} | {str(t[4]):<4} | {t[5]}")
            else:
                self.log.append("  （无属性）")

            # 方法任务
            self.log.append(f"\n📝 方法定义:")
            action_tasks = []
            for i, item in enumerate(actions):
                svc = match_service(item, services)
                siid = svc["siid"] if svc else item.get("siid", "?")
                if target_siid > 0 and str(siid) != str(target_siid):
                    continue
                sname = svc.get("description", svc.get("name", "")) if svc else "❌ 未匹配"
                action_tasks.append((i+1, item.get("name","?"), siid, sname))
                plan_rows.append({
                    "type": "方法",
                    "index": i + 1,
                    "name": item.get("name", ""),
                    "description": item.get("description", ""),
                    "format": "",
                    "value_type": "",
                    "siid": siid,
                    "service": sname,
                    "plan_status": "待创建" if svc else "服务未匹配",
                    "note": "dry-run 未执行",
                })

            if action_tasks:
                self.log.append(f"{'#':>3} | {'name':<20} | siid | 服务")
                self.log.append("-" * 60)
                for t in action_tasks:
                    self.log.append(f"{t[0]:>3} | {t[1]:<20} | {str(t[2]):<4} | {t[3]}")
            else:
                self.log.append("  （无方法）")

            # 事件任务
            self.log.append(f"\n📝 事件定义:")
            event_tasks = []
            for i, item in enumerate(events):
                svc = match_service(item, services)
                siid = svc["siid"] if svc else item.get("siid", "?")
                if target_siid > 0 and str(siid) != str(target_siid):
                    continue
                sname = svc.get("description", svc.get("name", "")) if svc else "❌ 未匹配"
                event_tasks.append((i+1, item.get("name","?"), siid, sname))
                plan_rows.append({
                    "type": "事件",
                    "index": i + 1,
                    "name": item.get("name", ""),
                    "description": item.get("description", ""),
                    "format": "",
                    "value_type": "",
                    "siid": siid,
                    "service": sname,
                    "plan_status": "待创建" if svc else "服务未匹配",
                    "note": "dry-run 未执行",
                })

            if event_tasks:
                self.log.append(f"{'#':>3} | {'name':<20} | siid | 服务")
                self.log.append("-" * 60)
                for t in event_tasks:
                    self.log.append(f"{t[0]:>3} | {t[1]:<20} | {str(t[2]):<4} | {t[3]}")
            else:
                self.log.append("  （无事件）")

            total = len(prop_tasks) + len(action_tasks) + len(event_tasks)
            self.log.append(f"\n🏁 共 {total} 条（属性{len(prop_tasks)}+方法{len(action_tasks)}+事件{len(event_tasks)}，干跑未执行）")
            plan_path = default_desktop_path("miot_dry_run_plan.xlsx")
            write_dry_run_plan(plan_path, plan_rows)
            self.log.append(f"📋 dry-run 计划表已保存: {plan_path}")
        except Exception:
            self.log.append(f"❌ {traceback.format_exc()}")

    def _start_create(self):
        config, items = self._load()
        if not config:
            return
        props, actions, events = items

        # ─── 优先检查产品状态 ──────────────────────────────────
        self.log.clear()
        self.log.append("🔍 正在检查产品状态...")
        try:
            is_ok, status, status_name, msg = check_product_status(config)
            if is_ok:
                self.log.append(f"✅ {msg}")
            else:
                self.log.append(f"❌ {msg}")
                QMessageBox.critical(self, "产品状态检查失败", msg)
                return
        except Exception as e:
            self.log.append(f"⚠️ 产品状态检查异常: {e}（继续执行）")

        total = len(props) + len(actions) + len(events)
        reply = QMessageBox.question(
            self, "确认创建",
            f"即将创建 {total} 条（属性{len(props)}+方法{len(actions)}+事件{len(events)}），是否继续？",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        if reply != QMessageBox.StandardButton.Yes:
            return
        try:
            services = create_query_services(config)
        except Exception as e:
            QMessageBox.critical(self, "错误", f"查询服务失败: {e}"); return

        target_siid = self.siid_spin.value()

        # 构建 task 列表：(type_label, item, build_fn, create_fn, id_field)
        all_tasks = []
        task_summaries = []
        for i, p in enumerate(props):
            svc = match_service(p, services)
            siid = svc["siid"] if svc else p.get("siid", "?")
            if target_siid > 0 and str(siid) != str(target_siid):
                continue
            all_tasks.append(("属性", p, build_request_body, create_property, "piid", svc))
            task_summaries.append(("属性定义", {"index": i + 1, "name": p.get("name", ""), "siid": siid}))
        for i, a in enumerate(actions):
            svc = match_service(a, services)
            siid = svc["siid"] if svc else a.get("siid", "?")
            if target_siid > 0 and str(siid) != str(target_siid):
                continue
            all_tasks.append(("方法", a, build_action_request_body, create_action, "aiid", svc))
            task_summaries.append(("方法定义", {"index": i + 1, "name": a.get("name", ""), "siid": siid}))
        for i, e in enumerate(events):
            svc = match_service(e, services)
            siid = svc["siid"] if svc else e.get("siid", "?")
            if target_siid > 0 and str(siid) != str(target_siid):
                continue
            all_tasks.append(("事件", e, build_event_request_body, create_event, "eiid", svc))
            task_summaries.append(("事件定义", {"index": i + 1, "name": e.get("name", ""), "siid": siid}))

        if not all_tasks:
            QMessageBox.information(self, "提示", "没有匹配的任务"); return

        task_errors = []
        for label in ("属性定义", "方法定义", "事件定义"):
            task_errors.extend(validate_tasks(
                [task for task_label, task in task_summaries if task_label == label],
                label,
            ))
        if task_errors:
            msg = "\n".join(task_errors[:20])
            self.log.append("\n❌ 任务校验未通过:\n" + msg)
            QMessageBox.warning(self, "任务校验未通过", msg)
            return

        self.log.clear()
        self.log.append(f"🚀 开始创建 {len(all_tasks)} 条（属性+方法+事件）...\n")
        self._set_btns(running=True)
        self.progress.setVisible(True); self.progress.setRange(0, len(all_tasks))

        self._worker = CreateAllWorker(config, all_tasks, services, self.delay_spin.value() / 1000.0)
        self._worker.progress.connect(self.log.append)
        self._worker.update_progress.connect(lambda c, t: self.progress.setValue(c))
        self._worker.finished_ok.connect(self._done_ok)
        self._worker.finished_err.connect(self._done_err)
        self._worker.start()

    def _cancel(self):
        if self._worker and self._worker.isRunning():
            self._worker.cancel()
            self.log.append("⚠️ 取消请求已发送")
        self._set_btns(running=False)

    def _done_ok(self, success, failed):
        self._set_btns(running=False)
        self.log.append(f"\n{'='*50}\n📊 成功 {success} / 失败 {failed} / 共 {success+failed}")
        if failed:
            QMessageBox.warning(self, "完成", f"成功 {success}, 失败 {failed}")
        else:
            QMessageBox.information(self, "完成", f"🎉 全部 {success} 条创建成功！")

    def _done_err(self, msg):
        self._set_btns(running=False)
        self.log.append(f"\n❌ {msg}")
        QMessageBox.critical(self, "失败", msg)

    def _set_btns(self, running):
        for b in (self.btn_create, self.btn_dry, self.btn_list):
            b.setEnabled(not running)
        self.btn_cancel.setEnabled(running)
        self.progress.setVisible(running)


# ─── Tab: 生成模板 ────────────────────────────────────────────

class TemplatePropTab(QWidget):
    def __init__(self):
        super().__init__()
        self._build()

    def _build(self):
        layout = QVBoxLayout(self)
        grp = QGroupBox("生成空白属性 Excel 模板")
        form = QFormLayout()
        self.out_edit = QLineEdit("MIoT_属性创建模板.xlsx")
        btn_br = QPushButton("浏览...")
        btn_br.clicked.connect(self._browse)
        row = QHBoxLayout(); row.addWidget(self.out_edit); row.addWidget(btn_br)
        form.addRow("输出路径:", row)
        btn_gen = QPushButton("📄 生成模板")
        btn_gen.setObjectName("successBtn")
        btn_gen.clicked.connect(self._gen)
        form.addRow("", btn_gen)
        grp.setLayout(form)
        layout.addWidget(grp)
        layout.addStretch()

    def _browse(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "选择输出路径", "MIoT_属性创建模板.xlsx", "Excel (*.xlsx)")
        if path:
            self.out_edit.setText(path)

    def _gen(self):
        path = self.out_edit.text().strip()
        if not path:
            QMessageBox.warning(self, "提示", "请填写输出路径"); return
        try:
            _generate_blank_template(path)
            QMessageBox.information(self, "成功", f"模板已生成:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "失败", str(e))


def _generate_blank_template(output_path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    header_font   = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill   = PatternFill("solid", fgColor="4472C4")
    opt_fill      = PatternFill("solid", fgColor="8DB4E2")
    header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border   = Border(left=Side(style="thin"), right=Side(style="thin"),
                           top=Side(style="thin"),  bottom=Side(style="thin"))
    desc_font     = Font(name="Arial", size=9, color="666666")
    desc_fill     = PatternFill("solid", fgColor="D9E2F3")
    opt_desc_fill = PatternFill("solid", fgColor="E8F0FE")

    ws = wb.active; ws.title = "属性定义"
    # 使用与 miot_common.PROPERTY_COLUMNS 一致的列定义（含 piid，列顺序一致）
    columns = [
        ("name",              20, "属性英文名\n如: on, mode, delay-time",          True),
        ("description",       25, "属性中文描述\n如: 开关, 模式, 延时时间",          True),
        ("format",            12, "数据格式\nbool/uint8/uint16/uint32\n/int8/int16/int32/float/string", True),
        ("service_desc",      22, "服务中文描述（推荐）\n如: 开关一键、按键1点动毫秒数", True),
        ("value_list",        35, "枚举值（仅enum类型）\n格式: 0:关闭,1:开启,2:待机",  False),
        ("value_range_min",   14, "数值最小值\n（仅number类型）",                    False),
        ("value_range_max",   14, "数值最大值\n（仅number类型）",                    False),
        ("value_range_step",  14, "数值步长\n（仅number类型）",                     False),
        ("service_name",      20, "服务英文名\n如: switch, jog-delay-time",         False),
        ("siid",               8, "服务ID（备选）\n直接指定siid，填了则忽略service匹配", False),
        ("access",            20, "访问权限\n默认: read,write,notify\n（gattAccess自动等同于access）", False),
        ("piid",               8, "属性ID\n（导出时自动填入，创建后可校验修正）", False),
    ]
    for i, (col, width, desc, required) in enumerate(columns, 1):
        cl = chr(64 + i)
        ws.column_dimensions[cl].width = width
        cell = ws.cell(row=1, column=i, value=col)
        cell.font = header_font
        cell.fill = header_fill if required else opt_fill
        cell.alignment = header_align; cell.border = thin_border
        dc = ws.cell(row=2, column=i, value=desc)
        dc.font = desc_font
        dc.fill = desc_fill if required else opt_desc_fill
        dc.alignment = Alignment(vertical="center", wrap_text=True)
        dc.border = thin_border
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 50
    dv = DataValidation(type="list", formula1='"bool,uint8,uint16,uint32,int8,int16,int32,float,string"', allow_blank=True)
    ws.add_data_validation(dv); dv.add("C3:C1000")

    ws2 = wb.create_sheet("公共配置")
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 65
    ws2.column_dimensions["C"].width = 40
    config_items = [
        ("template_version", TEMPLATE_VERSION, "模板版本（自动生成，请勿修改）", False),
        ("userId",       "", "小米账号用户ID（必填）",                    True),
        ("pdId",         "", "产品ID（必填）",                            True),
        ("model",        "", "设备型号（必填）",                           True),
        ("serviceToken", "", "浏览器 Cookie 获取（必填）",                 True),
        ("xiaomiiot_ph", "", "浏览器 Cookie 获取（必填）",                 True),
        ("connectType",  "16", "连接类型（默认16）",                       False),
        ("language",     "zh_cn", "语言（默认zh_cn）",                    False),
        ("version",      "1", "版本（默认1）",                            False),
        ("status",       "0", "状态（默认0）",                            False),
        ("source",       "4", "来源（默认4）",                            False),
        ("standard",     "false", "标准属性（默认false）",                 False),
        ("access",       "read,write,notify", "默认访问权限",              False),
    ]
    for i, (k, v, d, req) in enumerate(config_items, 1):
        kc = ws2.cell(row=i, column=1, value=k)
        ws2.cell(row=i, column=2, value=v)
        dc = ws2.cell(row=i, column=3, value=d)
        if req:
            kc.font = Font(name="Arial", bold=True, color="CC0000")
        dc.font = desc_font

    ws3 = wb.create_sheet("填写说明")
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 80
    instructions = [
        ("必填列", "name / description / format / service_desc"),
        ("枚举属性", "value_list 列填写格式: 0:关闭,1:开启,2:待机"),
        ("数值属性", "value_range_min / max / step 三列"),
        ("bool 属性", "format 填 bool，value_list 和 value_range 都留空"),
        ("服务匹配", "优先用 service_desc（服务中文名）匹配"),
        ("siid 列", "可选，填了则忽略 service 匹配"),
        ("access 列", "可选，默认 read,write,notify"),
    ]
    ws3.cell(row=1, column=1, value="项目").font = Font(bold=True, size=12)
    ws3.cell(row=1, column=2, value="说明").font = Font(bold=True, size=12)
    for i, (item, desc) in enumerate(instructions, 2):
        ws3.cell(row=i, column=1, value=item).font = Font(bold=True)
        ws3.cell(row=i, column=2, value=desc)

    wb.save(output_path)

